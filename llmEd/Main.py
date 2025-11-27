from pptx import Presentation

import Utility
#Inserimento file predefiniti .PPTX e JSON
layoutPPTX = 'LayoutLezioneSimulata_bis.pptx'
jsonFile = '/Users/alexandrasheykina/PycharmProjects/llmEd/documenti/olga.rossi/progettazione.json'
outputPPTX = 'presentazione.pptx'
formattingData = Utility.loadJsonData('formatting.json')
'''#Inserimento file predefiniti .PPTX e JSON
layoutPPTX = 'LayoutBacheche.pptx'
jsonFile = '/Users/alexandrasheykina/PycharmProjects/llmEd/documenti/sasa.sheykina/output_bacheche.json'
outputPPTX = 'presentazioneBacheche.pptx'
'''

#Stampa di generazione
print("Generazione in corso... Attendere!!")

#Caricamento file .JSON da considerare
jsonData = Utility.loadJsonData(jsonFile)

#Invocazione metodo main
Utility.replaceTag(layoutPPTX, jsonData, formattingData, outputPPTX)

#Stampa di successo
print("Presentazione generata con successo")
'''
prs = Presentation(layoutPPTX)
for slide_index, slide in enumerate(prs.slides):
    print(f"\n--- Slide {slide_index + 1} ---")
    for shape_index, shape in enumerate(slide.shapes):
        print(f"Shape {shape_index} (Index: {shape_index}):")
        print(f"  Type: {type(shape)}")
        print(f"  Name: {shape.name}")

        if shape.has_text_frame:
            print("  Text:")
            for paragraph in shape.text_frame.paragraphs:
                print(f"    - {paragraph.text}")
        elif shape.shape_type == 13:  # Picture type
            print("  This is a picture.")
        elif shape.shape_type == 19:  # Table type
            print("  This is a table.")
        else:
            print("  No recognizable content.")
            '''