import uuid

from docx.shared import RGBColor
import re
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.dml import MSO_COLOR_TYPE
import streamlit as st
import os
import copy

def contains_tags(text):
    # Usa una regex per verificare se ci sono tag del tipo {{{...}}}
    return bool(re.search(r'\{\{\{.*?\}\}\}', text))

def clenup_lezioneSimulata(folder_output):
    prs = Presentation(os.path.join(folder_output, "lezioneSimulata.pptx"))
    from pptx.opc.packuri import PackURI
    slides_to_remove = []

    # Trova le diapositive da rimuovere
    slides_to_remove = []

    for i, slide in enumerate(prs.slides):
        # Verifica se la diapositiva contiene uno dei tag
        slide_contains_tag = False

        # Itera attraverso tutte le forme di testo nella diapositiva
        for shape in slide.shapes:
            if hasattr(shape, 'text'):  # Verifica se la forma ha del testo
                if contains_tags(shape.text):  # Se trova un tag, segnala la diapositiva
                    slide_contains_tag = True
                    break

        if slide_contains_tag:
            slides_to_remove.append(i)

    # Rimuove le diapositive trovate, partendo dall'ultima
    for slide_index in reversed(slides_to_remove):
        xml_slides = prs.slides._sldIdLst
        del xml_slides[slide_index]  # Rimuove la diapositiva usando l'indice

    # Salva la presentazione modificata
    prs.save(os.path.join(st.session_state["user_dir"], "lezioneSimulata.pptx"))

def cleanup_presentation(prs):
    from pptx.opc.packuri import PackURI
    for i, slide in enumerate(prs.slides, 1):
        new_partname = PackURI(f'/ppt/slides/slide{i}.xml')
        slide.part.partname = new_partname
    return prs

def _save_font_configuration(font, doPrint=False):
    saved = {}
    saved['name'] = font.name
    saved['size'] = font.size
    saved['bold'] = font.bold
    saved['italic'] = font.italic
    saved['underline'] = font.underline
    saved['color.type'] = font.color.type
    if doPrint:
        print(font.name, font.size)
    if font.color.type == MSO_COLOR_TYPE.SCHEME:
        saved['color.brightness'] = font.color.brightness
        saved['color.theme_color'] = font.color.theme_color
    elif font.color.type == MSO_COLOR_TYPE.RGB:
        saved['color.rgb'] = None if font.color.rgb is None else str(font.color.rgb)
        # saved['fill'] = font.fill
        # saved['language_id'] = font.language_id
    return saved

def _restore_font_configuration(saved, font):
    font.name = saved['name']
    font.size = saved['size']
    font.bold = saved['bold']
    font.italic = saved['italic']
    font.underline = saved['underline']
    if saved['color.type'] == MSO_COLOR_TYPE.SCHEME:
        font.color.brightness = saved['color.brightness']
        font.color.theme_color = saved['color.theme_color']
    elif saved['color.type'] == MSO_COLOR_TYPE.RGB:
        if saved['color.rgb'] is not None:
            font.color.rgb = RGBColor.from_string(saved['color.rgb'])
        else:
            font.color.rgb = None


def remove_all_shapes_from_slide(slide):
    # Ottieni il riferimento alla lista delle forme nel livello XML
    spTree = slide.shapes._spTree
    
    # Crea una lista degli elementi da rimuovere
    elements_to_remove = [shape.element for shape in slide.shapes]
    
    # Rimuovi ciascun elemento dalla lista
    for element in elements_to_remove:
        spTree.remove(element)


def copy_slide_background(src_slide, curr_slide):
    if src_slide.background and src_slide.background.fill.type == 'picture':
        curr_slide.background.fill.picture.image = src_slide.background.fill.picture.image


def copy_slide_elements(src_slide, curr_slide):
    # create images dict
    imgDict = {}

    remove_all_shapes_from_slide(curr_slide)
    # Copia lo sfondo se presente
    copy_slide_background(src_slide, curr_slide)
    # now copy contents from external slide, but do not copy slide properties
    # e.g. slide layouts, etc., because these would produce errors, as diplicate
    # entries might be generated
    for shp in src_slide.shapes:
        if shp.shape_type == 13:  # 13 corrisponde al tipo immagine
            # Genera un nome unico per l'immagine
            ext = shp.image.ext
            unique_name = f"{uuid.uuid4()}.{ext}"

            # Salva l'immagine su disco temporaneamente
            with open(unique_name, 'wb') as f:
                f.write(shp.image.blob)

            # Aggiungi l'immagine al dizionario con le sue proprietà
            imgDict[unique_name] = [shp.left, shp.top, shp.width, shp.height]
        else:
            # Copia altre forme
            el = shp.element
            newel = copy.deepcopy(el)
            curr_slide.shapes._spTree.insert_element_before(newel, 'p:extLst')

        # Aggiungi immagini alla diapositiva
    for img_name, props in imgDict.items():
        curr_slide.shapes.add_picture(img_name, props[0], props[1], props[2], props[3])
        os.remove(img_name)  # Rimuovi il file immagine temporaneo

def insert_slides_at_tag(main_pptx, additional_pptx, output_pptx, tag):
    main_prs = Presentation(main_pptx)
    additional_prs = Presentation(additional_pptx)
    
    tag_slide_index = None

    # Trova l'indice della diapositiva contenente il tag
    for slide_index, slide in enumerate(main_prs.slides):
        if any(shape.has_text_frame and tag in shape.text for shape in slide.shapes):
            tag_slide_index = slide_index
            break

    if tag_slide_index is not None:
        # Inserisci nuove diapositive mantenendo l'ordine originale
        new_slide_ids = []
        # Aggiungi le nuove diapositive
        for additional_slide in additional_prs.slides:
            idx = next(
                (i for i, layout in enumerate(main_prs.slide_layouts) if
                 layout.name == additional_slide.slide_layout.name),
                0
            )
            target_slide = main_prs.slides.add_slide(main_prs.slide_layouts[idx])
            copy_slide_elements(additional_slide, target_slide)
            new_slide_ids.append(main_prs.slides._sldIdLst[-1])  # Registra il nuovo ID

            # Inserisci i nuovi ID subito dopo il tag, mantenendo l'ordine originale
        tag_slide_position = list(main_prs.slides._sldIdLst).index(main_prs.slides._sldIdLst[tag_slide_index])
        for new_id in reversed(new_slide_ids):  # Inserisci in ordine inverso per mantenere l'ordine corretto
            main_prs.slides._sldIdLst.insert(tag_slide_position + 1, new_id)

        # Rimuovi la diapositiva contenente il tag
        xml_slides = main_prs.slides._sldIdLst
        slide_to_remove = xml_slides[tag_slide_index]
        rId = slide_to_remove.rId
        del xml_slides[tag_slide_index]  # Elimina la diapositiva dall'elenco
        main_prs.part.drop_rel(rId)  # Elimina la relazione corrispondente

    # Salva la presentazione finale
    cleanup_presentation(main_prs)
    main_prs.save(output_pptx)
    print(f"Slides inserite e slide con il tag '{tag}' rimossa.")


def remove_slide_with_tag(pptx_path, output_pptx, tag):
    prs = Presentation(pptx_path)

    tag_slide_index = None
    for slide_index, slide in enumerate(prs.slides):
        if any(shape.has_text_frame and tag in shape.text for shape in slide.shapes):
            tag_slide_index = slide_index
            break

    if tag_slide_index is not None:
        xml_slides = prs.slides._sldIdLst
        slide_to_remove = xml_slides[tag_slide_index]
        rId = slide_to_remove.rId
        del xml_slides[tag_slide_index]  # Rimuove la diapositiva
        prs.part.drop_rel(rId)  # Elimina la relazione corrispondente

        prs.save(output_pptx)
        print(f"Diaspositiva con tag '{tag}' rimossa con successo!")
    else:
        print(f"Nessuna diapositiva trovata con il tag '{tag}'.")

'''if __name__ == "__main__":
    #insert_slides_at_tag("lezioneSimulata.pptx", "verificaPrerequisiti.pptx", "lezioneSimulata.pptx",
    #                     "{{{verificaPrerequisiti.pptx}}}")
    insert_slides_at_tag("lezioneSimulata.pptx", "nucleoCentrale.pptx", "lezioneSimulata.pptx",
                         "{{{nucleoCentrale.pptx}}}")
    
    main_pptx = "presentazione.pptx"
    additional_pptx = "presentazioneBacheche.pptx"
    output_pptx = "final_pres.pptx"
    tag = "{{{presentazioneBacheche.pptx}}}"

    insert_slides_at_tag(main_pptx, additional_pptx, output_pptx, tag)'''