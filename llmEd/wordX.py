import json
from docx import Document
from docx.shared import Pt
import re
import copy

# Funzione per caricare i dati dal file JSON
def load_json_data(json_file):
    with open(json_file, 'r', encoding='utf-8') as file:
        data = json.load(file)
        print("LOAD JSON DATA")
        print(data)
    return data




# Funzione per ottenere il valore da una chiave nel JSON
def get_value_from_json(data, key_path):
    keys = key_path.split('.')
    current_data = data

    for key in keys:
        if isinstance(current_data, list):
            try:
                index = int(key)
                current_data = current_data[index]
            except (ValueError, IndexError):
                return None
        elif isinstance(current_data, dict):
            current_data = current_data.get(key, None)
        else:
            return None

    return current_data



# Funzione per sostituire i tag nel file Word
def replace_placeholder_in_paragraph(layout_docx, json_data, output_docx):
    doc = Document(layout_docx)
    pattern = r"<#(.*?)#>"
    for paragraph in doc.paragraphs:
        text =paragraph.text.strip()
        matches = re.findall(pattern, text)
        for match in matches:
            # Dividi il tag in parti
            main_key, *sub_keys = match.split("->")
            # Estrai il valore principale
            value = get_value_from_json(json_data, main_key)

            # Se ci sono sotto-chiavi, continua a navigare
            if sub_keys and isinstance(value, dict):
                sub_key_path = "->".join(sub_keys)
                value = get_value_from_json(value, sub_key_path)
            # Sostituisci il testo nel documento
            paragraph.text = paragraph.text.replace(f"<#{match}#>", str(value))

    uda_list = json_data["Unita Didattiche"]  # Estrai la lista di Unità Didattiche

    # Trova la tabella modello
    template_table = doc.tables[0]

    def populate_table(table, json_data):
        pattern = r"<#(.*?)#>"  # Cerca i segnaposto nel formato <#...#>

        for row in table.rows:
            for cell in row.cells:
                text = cell.text
                matches = re.findall(pattern, text)  # Trova i segnaposto nel testo della cella
                print(text)
                for match in matches:
                    # Divide il percorso del tag in parti per gestire le chiavi
                    keys = match.split("->")

                    # Naviga nel dizionario JSON usando le chiavi
                    value = json_data
                    for key in keys:
                        if isinstance(value, dict):
                            value = value.get(key, None)
                        else:
                            value = None
                            break
                            # Speciale: Stampa competenze e livelli
                    if match == "CP->criteri_valutazione->Competenze" and isinstance(value, list):
                                value = ""
                                for competenza in json_data["CP"]["criteri_valutazione"][
                                    "Competenze"]:
                                    value += f"Competenza: {competenza['indicatore']}\n"
                                    value += f"Base:  {competenza['Base']}\n"
                                    value += f"Intermedio:  {competenza['Intermedio']}\n"
                                    value += f"Avanzato:  {competenza['Avanzato']}\n"
                                    value += "\n"
                                    print(value)
                    elif isinstance(value, list):
                        if match == "CP->relazione_individuale->domande":
                            print("entrato nel relazione individuale")
                            value = ""
                            for relazione in json_data["CP"]["relazione_individuale"]['domande']:
                                value += f"• {relazione}\n\n\n\n"

                        else:
                            value = "\n".join(f"• {item}" for item in value)

                    elif value is None:  # Se il valore non esiste
                        value = "N/A"
                    else:
                        value = str(value)

                    # Sostituisci il segnaposto nel testo della cella
                    cell.text = text.replace(f"<#{match}#>", value)

                # Modifica il font e la dimensione del testo nella cella
                for paragrafo in cell.paragraphs:
                    for run in paragrafo.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(12)

    # Rimuovere la tabella modello solo alla fine
    uda_tables = []

    for i, uda in enumerate(uda_list, start=2):
        new_table = copy.deepcopy(template_table)  # Crea una copia della tabella
        populate_table(new_table, uda)  # Riempie la tabella con i dati

        # Aggiungi il paragrafo "UNITÀ DIDATTICA" prima della tabella
        paragraph = doc.add_paragraph(f"UNITA’ DI APPRENDIMENTO {i}")
        paragraph_format = paragraph.paragraph_format
        paragraph_format.space_before = Pt(12)
        paragraph_format.space_after = Pt(12)
        run = paragraph.runs[0]
        run.font.name = "Times New Roman"
        run.font.size = Pt(16)
        run.bold = True

        # Aggiungi il paragrafo alla lista
        uda_tables.append((paragraph, new_table))

    # Inserisci in ordine corretto
    for paragraph, new_table in reversed(uda_tables):
        template_table._element.addnext(new_table._element)
        new_table._element.addnext(paragraph._element)

    # Rimuovi la tabella modello originale
    template_table._element.getparent().remove(template_table._element)

    last_paragraph, last_table = uda_tables[-1]
    last_paragraph._element.getparent().remove(last_paragraph._element)

    # Salva il documento aggiornato
    doc.save(output_docx)

'''if __name__ == "__main__":
    template_pptx = "LayoutProgettazioneDidattica.docx"
    json_file = "progettazione_didattica.json"
    output_pptx = "progettazione_didattica.docx"

    try:
        with open(json_file, 'r') as f_json:
            json_data = json.load(f_json)  # Correctly load JSON data from file
            print(json_data)
    except FileNotFoundError:
        print(f"Error: File {json_file} not found.")
        exit(1)  # Exit if the JSON file cannot be found
    except json.JSONDecodeError as e:
        print(f"Error decoding JSON from {json_file}: {e}")
        exit(1)  # Exit if there's a JSON parsing error

    replace_placeholder_in_paragraph(template_pptx, json_data, output_pptx)'''